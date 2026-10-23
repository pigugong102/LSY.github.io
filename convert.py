# coding=utf-8
import os
import html
import datetime
import openpyxl
from openpyxl.utils import get_column_letter

def rgb_to_hex(rgb_str):
    if not rgb_str or str(rgb_str) == '00000000' or "Values must be of type" in str(rgb_str):
        return None
    rgb_str = str(rgb_str)
    if len(rgb_str) == 8:
        return '#' + rgb_str[2:]
    elif len(rgb_str) == 6:
        return '#' + rgb_str
    return None

def is_dark_color(hex_color):
    if not hex_color or not hex_color.startswith('#'):
        return False
    h = hex_color.lstrip('#')
    if len(h) != 6:
        return False
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255.0
        return luminance < 0.62
    except ValueError:
        return False

def format_cell_value(val):
    if val is None:
        return ''
    if isinstance(val, (datetime.datetime, datetime.date)):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    s = str(val)
    if s.endswith(' 00:00:00'):
        s = s[:-9]
    return html.escape(s).replace('\n', '<br>')

def clean_sheet_title(sheet_name, file_name):
    if sheet_name in ['시트1', 'Sheet1']:
        return file_name.replace('.xlsx', '').replace('_', ' ')
    mapping = {
        '요구사항정의서': '요구사항 정의서',
        '요구사항_개정이력': '요구사항 개정이력',
        '기능정의서': '기능 정의서',
    }
    if sheet_name in mapping:
        return mapping[sheet_name]
    return sheet_name.replace('_', ' ')

def render_table_block(ws, min_row, max_row, min_col, max_col, merged_cells_map, project, file_name, is_subtable=False):
    table_style = ' style="margin-bottom: 24px;"' if is_subtable else ''
    out = f'<table{table_style}>\n<colgroup>\n'
    for c in range(min_col, max_col + 1):
        col_letter = get_column_letter(c)
        width = ws.column_dimensions[col_letter].width
        if project == 'planslot' and file_name == '테스트케이스.xlsx' and c == max_col:
            width = 50
        if width:
            px = min(int(width * 8), 520)
            out += f'<col style="width: {px}px;">\n'
        else:
            out += '<col>\n'
    out += '</colgroup>\n'

    for row in range(min_row, max_row + 1):
        out += '<tr>\n'
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            merge_info = merged_cells_map.get((row, col))
            if merge_info == 'skip':
                continue

            rowspan = merge_info['rowspan'] if merge_info else 1
            colspan = merge_info['colspan'] if merge_info else 1
            if row + rowspan - 1 > max_row:
                rowspan = max_row - row + 1
            if col + colspan - 1 > max_col:
                colspan = max_col - col + 1

            bg_color = None
            if cell.fill and cell.fill.start_color and getattr(cell.fill.start_color, 'rgb', None):
                bg_color = rgb_to_hex(cell.fill.start_color.rgb)

            # In aimatch WBS.xlsx, #F3F4F6 on non-status columns was manual row striping that got partially erased in Excel;
            # skip #F3F4F6 on non-status columns so CSS tr:nth-child(even) stripes rows cleanly without half-gray rows.
            if project == 'aimatch' and file_name == 'WBS.xlsx' and bg_color and bg_color.upper() == '#F3F4F6' and col != 9:
                bg_color = None

            styles = []
            if bg_color and bg_color.upper() not in ['#000000', '#FFFFFF']:
                styles.append(f'background-color: {bg_color}')
                if is_dark_color(bg_color):
                    styles.append('color: #FFFFFF')
                    styles.append('font-weight: 600')

            is_bold = bool(cell.font and cell.font.bold)
            if is_bold and 'font-weight: 600' not in styles:
                styles.append('font-weight: 600')

            style_attr = f' style="{"; ".join(styles)};"' if styles else ''
            attrs = f' rowspan="{rowspan}" colspan="{colspan}"{style_attr}'

            val = format_cell_value(cell.value)
            if is_bold and val:
                val = f'<b>{val}</b>'

            out += f'<td{attrs}>{val}</td>\n'
        out += '</tr>\n'
    out += '</table>\n'
    return out

def excel_to_html(excel_path, html_path, file_name, project):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    html_content = '''<html><head><meta charset="utf-8">
    <style>
    body { font-family:-apple-system, BlinkMacSystemFont, "Pretendard Variable", Pretendard, sans-serif; padding:24px; margin:0; background:#f9f9fa; }
    h2 { color:#111; margin:32px 0 16px; font-size:20px; border-bottom:2px solid #222; padding-bottom:8px; display:inline-block; }
    h2:first-of-type { margin-top: 8px; }
    table { table-layout:auto; border-collapse:collapse; width:100%; max-width:1400px; font-size:13.5px; text-align:left; background:#fff; margin-bottom:48px; box-shadow:0 2px 8px rgba(0,0,0,0.04); border-radius:6px; overflow:hidden; }
    th, td { border:1px solid #e0e0e0; padding:10px 14px; word-break:keep-all; vertical-align:middle; line-height:1.6; color:#333; }
    tr:nth-child(even) { background:#fafafa; }
    </style></head><body>'''

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        non_empty_rows = []
        non_empty_cols = []
        rows_with_val = set()
        for row in ws.iter_rows():
            for cell in row:
                has_val = cell.value is not None and str(cell.value).strip() != ''
                bg = rgb_to_hex(cell.fill.start_color.rgb) if (cell.fill and cell.fill.start_color and getattr(cell.fill.start_color, 'rgb', None)) else None
                has_bg = bg is not None and bg.upper() not in ['#000000', '#FFFFFF']
                if has_val:
                    rows_with_val.add(cell.row)
                if has_val or has_bg:
                    non_empty_rows.append(cell.row)
                    non_empty_cols.append(cell.column)

        # Skip empty sheets or sheets that only have a single header row and no data rows
        if not non_empty_rows or len(rows_with_val) <= 1:
            continue

        min_row, max_row = min(non_empty_rows), max(non_empty_rows)
        min_col, max_col = min(non_empty_cols), max(non_empty_cols)

        for merged_range in ws.merged_cells.ranges:
            mc_min_col, mc_min_row, mc_max_col, mc_max_row = merged_range.bounds
            if mc_min_row >= min_row and mc_max_row > max_row:
                max_row = mc_max_row
            if mc_min_col >= min_col and mc_max_col > max_col:
                max_col = mc_max_col

        if project == 'planslot' and file_name == 'WBS.xlsx':
            max_row -= 4
            max_col -= 1

        merged_cells_map = {}
        for merged_range in ws.merged_cells.ranges:
            mc_min_col, mc_min_row, mc_max_col, mc_max_row = merged_range.bounds
            merged_cells_map[(mc_min_row, mc_min_col)] = {
                'rowspan': mc_max_row - mc_min_row + 1,
                'colspan': mc_max_col - mc_min_col + 1
            }
            for r in range(mc_min_row, mc_max_row + 1):
                for c in range(mc_min_col, mc_max_col + 1):
                    if (r, c) != (mc_min_row, mc_min_col):
                        merged_cells_map[(r, c)] = 'skip'

        display_sheet_name = clean_sheet_title(sheet_name, file_name)
        html_content += f'<h2>{html.escape(display_sheet_name)}</h2>\n'

        # Special handling for aimatch WBS.xlsx: separate main WBS (rows 1..45) and compact Legend (rows 46..47, cols 1..10)
        if project == 'aimatch' and file_name == 'WBS.xlsx':
            html_content += render_table_block(ws, 1, 45, 1, 18, merged_cells_map, project, file_name, is_subtable=True)
            legend_map = dict(merged_cells_map)
            legend_map[(46, 1)] = {'rowspan': 1, 'colspan': 10}
            for c in range(2, 11):
                legend_map[(46, c)] = 'skip'
            html_content += '<table style="max-width: 680px; margin-bottom: 48px;">\n'
            for r in [46, 47]:
                html_content += '<tr>\n'
                for c in range(1, 11):
                    cell = ws.cell(r, c)
                    mi = legend_map.get((r, c))
                    if mi == 'skip':
                        continue
                    rs = mi['rowspan'] if mi else 1
                    cs = mi['colspan'] if mi else 1
                    bg_color = rgb_to_hex(cell.fill.start_color.rgb) if (cell.fill and cell.fill.start_color and getattr(cell.fill.start_color, 'rgb', None)) else None
                    styles = []
                    if r == 46:
                        styles.extend(['background-color: #1B3A6B', 'color: #FFFFFF', 'font-weight: 600'])
                    elif bg_color and bg_color.upper() not in ['#000000', '#FFFFFF']:
                        styles.append(f'background-color: {bg_color}')
                        styles.append('width: 32px')
                    if cell.font and cell.font.bold and 'font-weight: 600' not in styles:
                        styles.append('font-weight: 600')
                    style_attr = f' style="{"; ".join(styles)};"' if styles else ''
                    val = format_cell_value(cell.value)
                    if cell.font and cell.font.bold and val:
                        val = f'<b>{val}</b>'
                    html_content += f'<td rowspan="{rs}" colspan="{cs}"{style_attr}>{val}</td>\n'
                html_content += '</tr>\n'
            html_content += '</table>\n'
            continue

        # Special handling for planslot 클래스정의서.xlsx: split main table (cols 1..7) and side response table (cols 9..10, rows 1..5)
        if project == 'planslot' and file_name == '클래스정의서.xlsx':
            html_content += render_table_block(ws, 1, max_row, 1, 7, merged_cells_map, project, file_name, is_subtable=True)
            html_content += '<table style="max-width: 420px; margin-bottom: 48px;">\n'
            html_content += render_table_block(ws, 1, 5, 9, 10, merged_cells_map, project, file_name, is_subtable=False)[len('<table>\n'):]
            continue

        def is_empty_separator_row(r):
            for c in range(min_col, max_col + 1):
                if merged_cells_map.get((r, c)) is not None:
                    return False
                cell = ws.cell(r, c)
                has_val = cell.value is not None and str(cell.value).strip() != ''
                bg = rgb_to_hex(cell.fill.start_color.rgb) if (cell.fill and cell.fill.start_color and getattr(cell.fill.start_color, 'rgb', None)) else None
                has_bg = bg is not None and bg.upper() not in ['#000000', '#FFFFFF']
                if has_val or has_bg:
                    return False
            return True

        blocks = []
        curr_start = None
        for r in range(min_row, max_row + 1):
            if is_empty_separator_row(r):
                if curr_start is not None:
                    blocks.append((curr_start, r - 1))
                    curr_start = None
            else:
                if curr_start is None:
                    curr_start = r
        if curr_start is not None:
            blocks.append((curr_start, max_row))

        for idx, (b_start, b_end) in enumerate(blocks):
            is_sub = (len(blocks) > 1 and idx < len(blocks) - 1)
            html_content += render_table_block(ws, b_start, b_end, min_col, max_col, merged_cells_map, project, file_name, is_subtable=is_sub)

    html_content += '''
    <script>
    document.addEventListener('keydown', function(e) {
      if (e.key === 'Escape') {
        if (window.parent && window.parent.closeDocModal) {
          window.parent.closeDocModal();
        }
      }
    });
    </script>
    </body></html>'''
    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html_content)

for project in ['planslot', 'aimatch']:
    docs_path = f'docs/{project}'
    if not os.path.exists(docs_path):
        continue
    for file in sorted(os.listdir(docs_path)):
        if file.endswith('.xlsx'):
            try:
                excel_to_html(
                    os.path.join(docs_path, file),
                    os.path.join(docs_path, file.replace('.xlsx', '.html')),
                    file,
                    project
                )
                print(f'Converted {project}/{file}')
            except Exception as e:
                print(f'Error converting {project}/{file}: {e}')
